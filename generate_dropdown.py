import csv
from collections import defaultdict
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

# The provided list of all companies
company_list_raw = """한국전자통신연구원
한국화학연구원
한국에너지기술연구원
한국지질자원연구원
한국생명공학연구원
한국원자력연구원
한국표준과학연구원
한국기계연구원
한국기초과학지원연구원 부설 국가핵융합연구소
한국기초과학지원연구원
한국항공우주연구원
한국과학기술정보연구원
한국화학연구원 부설 안전성평가연구소
한국해양과학기술원 부설 선박해양플랜트 연구소
한국한의학연구원
한국천문연구원
한국전자통신연구원 부설 국가보안기술연구소
기초과학연구원 부설 국가수리과학연구소
기초과학연구원
한국과학기술원 부설 나노종합기술원
한국연구재단
한국원자력안전기술원
정보통신산업진흥원 부설 정보통신기술진흥센터
한국원자력통제기술원
한국산업기술시험원 중부지역본부
국방과학연구소
배재대학교
과학기술연합대학원대학교(UST)
충남대학교
한국과학기술원
대덕대학교
한남대학교
한밭대학교
(재)다차원스마트 IT융합시스템연구단
대전교육과학연구원
한국전력공사 전력연구원
K-water 연구원
교통안전공단 중부지역본부
안전보건공단 산업안전보건연구원 산업화학연구실(화학물질독성연구실)
한국건설생활환경시험연구원 대전충남지원
한국토지주택공사 토지주택연구원
한국수력원자력㈜ 중앙연구원
국제지식재산연수원
금강유역환경청
국립중앙과학관
대전인재개발원
대전교육정보원
대전지방기상청
대전충남지방중소기업청
대전세관
정부통합전산센터
화학물질안전원
대전광역시 시설관리공단 무지개복지센터
한국전력기술㈜ 원자로설계개발단
연구개발특구진흥재단
(재)대전경제통상진흥원
대전직업능력개발원
한국원자력연료㈜
안전보건공단 대전지역본부
(재)대전테크노파크
한국조폐공사
한국원자력환경공단
대전시민천문대
중소기업기술정보진흥원
대전충남KOTRA지원단
대전신용보증재단 북부지점
중소기업진흥공단 청년사관학교
국가과학기술인력개발원 교육센터
창업진흥원
한국서부발전
대전창조경제혁신센터
한국수자원공사
정보통신기획평가원
한국철도공사
코레일테크㈜
한국철도시설공단
한국산림복지진흥원
한국가스기술공사
소상공인진흥공단
신용보증재단중앙회
목원대학교 산학협력단
(사)한국방사성폐기물학회
대덕이노폴리스 벤처협회
대전대학교 산학협력단
대덕산업단지관리공단
한국여성원자력전문인협회
대한산업보건협회 대전산업보건센터
(사)출연(연)연구발전협의회 총연합회
대한산업안전협회 대전지역본부
(사)대덕클럽
한국단미사료협회
(재)한국원자력협력재단
(재)장애인기업종합지원센터
대덕연구개발특구기관장협의회
(재)중앙문화재연구원 대전사무소
대전지역사업평가단
대전녹색환경지원센터
(사)대한여성과학기술인회
한국표준협회 대전세종충남지역본부
창업진흥원 창업보육센터
대덕원자력포럼
한국산업기술진흥협회
한국기술사업화진흥협회
한국여성경제인협회
대전문화산업진흥협회
한국과학기술단체총연합회 (과총)
(재)기가코리아사업단
나노융합산업연구조합
바이오메스연구단
(사)과학기술연우연합회
계룡건설산업
골프존
나노신소재
나노팀
나노엔텍
네오팜
노타
뉴로스
대산F&B
동양에스텍
디엔에프
라이온켐텍
라이트론
레인보우로보틱스
리가켐바이오
리메드
민테크
바이오니아
비비씨
비플라이소프트
빛과전자
수젠텍
시스웍
신테카바이오
쎄트렉아이
아이디스
아이디스홀딩스
아이비전웍스
아이쓰리시스템
아이빔테크놀로지
안지오랩
알루코
알테오젠
에르코스
에이치엔에스하이텍
엔지켐생명과학
오름테라퓨틱
와이바이오로직스
우성
원텍
위드텍
위월드
이비테크
인텍플러스
인투셀
잼코
제노텍
제일사료
젬백스
지노믹트리
진시스템
컨텍
크로우
큐로셀
토모큐브
파이버프로
펩트론
플라즈맵
프리시젼바이오
한독크린텍
한빛레이저
한온시스템
한켐
HLB제넥스
HLB파나진
KT&G
KTcs
LX세미콘"""

all_companies = [c.strip() for c in company_list_raw.splitlines() if c.strip()]

# Cleaned data logic
parsed_data = [
    ('3/23', '공공', '조연우', '대전광역시 외국인주민 통합지원센터'), # Not in list, but let's just insert it anyway or map it
    ('3/23', '상장', '조연우', '월드번역원((주)월시스)'),
    ('3/23', '공공', '허준혁', '정부통합전산센터'),
    ('3/23', '공공', '오채영', '소상공인진흥공단'),
    ('3/23', '상장', '오채영', '한독크린텍'),
    ('3/30', '공공', '구현경', '대전녹색환경지원센터'),
    ('3/30', '공공', '원창호', '한국항공우주연구원'),
    ('3/30', '공공', '백재현', '대전지방기상청'),
    ('3/30', '상장', '허준혁', '비플라이소프트'),
    ('3/30', '공공', '이승주', '한국수자원공사'),
    ('3/30', '상장', '강체첵', 'KT&G'),
    ('3/30', '상장', '진승훈', '큐로셀'),
    ('3/23', '상장', '전은재', '쎄트렉아이'),
    ('4/6', '상장', '구현경', 'KT&G'),
    ('4/6', '공공', '오다인', '대전직업능력개발원'),
    ('4/6', '공공', '박해원', '정보통신기획평가원'),
    ('4/6', '상장', '이승주', '계룡건설산업'),
    ('4/6', '상장', '백재현', '에르코스'),
    ('4/6', '공공', '이현기', '한국철도시설공단'),
    ('4/13', '공공', '양하영', '국방과학연구소'),
    ('4/13', '공공', '이수호', '한국전자통신연구원'),
    ('4/13', '상장', '정현준', '동양에스텍'),
    ('4/13', '상장', '원창호', '펩트론'),
    ('4/13', '상장', '안은률', '대산F&B'),
    ('4/13', '공공', '이준성', '중소기업기술정보진흥원'),
    ('4/6', '공공', '유태수', '코레일테크㈜'), # corrected name for matching
    ('4/13', '상장', '곽현지', '골프존'),
    ('4/20', '공공', '구민경', '대전시민천문대'),
    ('4/20', '공공', '신재민', '한국가스기술공사'),
    ('4/20', '공공', '진승훈', '신용보증재단중앙회'),
    ('5/4', '공공', '김채은', '대전세관'),
    ('5/4', '상장', '오다인', '우성'),
    ('5/4', '상장', '이수호', '인텍플러스'),
    ('5/4', '상장', '박해원', '아이디스'),
    ('5/4', '상장', '이준희', 'KTcs'),
    ('5/4', '상장', '정구영', '코셈'),
    ('5/11', '공공', '안은률', '대전문화산업진흥협회'),
    ('5/11', '상장', '김주헌', '아이디스홀딩스'),
    ('5/11', '공공', '김주헌', '안전보건공단대전지역본부'), # note: mapped slightly
    ('5/11', '상장', '이준성', '위드텍'),
    ('5/11', '상장', '유태수', '지노믹트리'),
    ('5/11', '공공', '한성경', '한국산림복지진흥원'),
    ('5/11', '공공', '이채원', '대전창조경제혁신센터'),
    ('5/18', '상장', '양하영', '알루코'),
    ('5/18', '공공', '정현준', '한국화학연구원'),
    ('5/18', '상장', '정하진', '원텍'),
    ('5/18', '공공', '정구영', '한국수자원공사'), # k-water
    ('5/18', '상장', '이채원', '네오팜'),
    ('5/18', '공공', '전은재', '한국기계연구원'),
    ('5/18', '공공', '곽현지', '(재)대전경제통상진흥원'),
    ('6/1', '상장', '구민경', '컨텍'),
    ('6/1', '공공', '양예지', '대전도시공사'), # not in list maybe
    ('6/1', '상장', '신재민', '파이버프로'),
    ('6/1', '상장', '양예지', '레인보우로보틱스')
]

dates = ['3/23', '3/30', '4/6', '4/13', '4/20', '5/4', '5/11', '5/18', '6/1']

# Group data
schedule_data = defaultdict(lambda: {'공공': [], '상장': []})
for pd, pg, pn, pc in parsed_data:
    schedule_data[pd][pg].append((pc, pn)) # Company, Name

wb = openpyxl.Workbook()
ws_list = wb.active
ws_list.title = "목록"
for i, c in enumerate(all_companies, 1):
    ws_list.cell(row=i, column=1, value=c)

# Create main sheet
ws = wb.create_sheet("발표일정", 0)

# Build Header
headers = ["날짜"]
for i in range(1, 5):
    headers.extend([f"공공기관 {i} 기업명", f"공공기관 {i} 이름"])
for i in range(1, 5):
    headers.extend([f"상장기업 {i} 기업명", f"상장기업 {i} 이름"])
ws.append(headers)

# Styling for header
header_fill_pub = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
header_fill_pri = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
for col in range(2, 10):
    ws.cell(row=1, column=col).fill = header_fill_pub
for col in range(10, 18):
    ws.cell(row=1, column=col).fill = header_fill_pri

# Adding data validation drop down
# It will apply to the columns corresponding to "기업명"
dv = DataValidation(type="list", formula1="=목록!$A$1:$A$" + str(len(all_companies)), allow_blank=True)
ws.add_data_validation(dv)

# Track duplicates
company_count = defaultdict(int)
for _, _, _, pc in parsed_data:
    # simple norm
    norm_c = str(pc).lower().replace(" ", "")
    if norm_c == "k-water": norm_c = "한국수자원공사"
    elif "코레일테크" in norm_c: norm_c = "코레일테크㈜"
    elif "경제통상" in norm_c: norm_c = "(재)대전경제통상진흥원"
    elif norm_c == "kt&g": norm_c = "KT&G"
    elif norm_c == "대산f&b": norm_c = "대산F&B"
    company_count[norm_c] += 1

red_font = Font(color="FF0000")

row_idx = 2
for d in dates:
    pub_list = schedule_data[d]['공공']
    pri_list = schedule_data[d]['상장']
    
    # We pad the lists up to 4 items
    while len(pub_list) < 4: pub_list.append(("", ""))
    while len(pri_list) < 4: pri_list.append(("", ""))
    
    row_data = [d]
    for c, n in pub_list[:4]: 
        row_data.extend([c, n])
    for c, n in pri_list[:4]:
        row_data.extend([c, n])
        
    ws.append(row_data)
    
    # Apply validations and find duplicates
    for col_idx in [2, 4, 6, 8, 10, 12, 14, 16]:
        cell = ws.cell(row=row_idx, column=col_idx)
        dv.add(cell)
        
        # Check for duplicates or special flags
        comp_name = cell.value
        if comp_name:
            norm_c = str(comp_name).lower().replace(" ", "")
            if norm_c == "k-water": norm_c = "한국수자원공사"
            elif "코레일테크" in norm_c: norm_c = "코레일테크㈜"
            elif "경제통상" in norm_c: norm_c = "(재)대전경제통상진흥원"
            elif norm_c == "kt&g": norm_c = "KT&G"
            elif norm_c == "대산f&b": norm_c = "대산F&B"
            
            # Find the name cell mapping to this
            name_cell = ws.cell(row=row_idx, column=col_idx+1)
            
            if company_count[norm_c] > 1:
                # Add note in the name cell
                name_cell.value = str(name_cell.value) + " (중복선정)"
                name_cell.font = red_font
            
            if name_cell.value and '유태수' in str(name_cell.value) and d == '4/6':
                name_cell.value = str(name_cell.value) + " (6주차/4.13작성건)"
                name_cell.font = red_font

    row_idx += 1

# Formatting
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=17):
    for cell in row:
        cell.alignment = Alignment(vertical='center', horizontal='center')
        
# Adjust widths
ws.column_dimensions['A'].width = 10
for i in [2, 4, 6, 8, 10, 12, 14, 16]:  # Company names
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 25
for i in [3, 5, 7, 9, 11, 13, 15, 17]:  # Names
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 25

# Save
wb.save('c:/Users/Sumin/00_Projects/지역사회의이해발표시트/발표일정_드롭다운.xlsx')
