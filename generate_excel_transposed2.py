import csv
from collections import defaultdict
import openpyxl
from openpyxl.styles import Alignment, Font
from copy import copy

parsed_data = [
    ('3/23', '공공', '조연우', '대전광역시 외국인주민 통합지원센터'),
    ('3/23', '상장', '조연우', '월드번역원((주)월시스)'),
    ('3/23', '공공', '허준혁', '정부통합전산센터'),
    ('3/23', '공공', '오채영', '소상공인진흥공단'),
    ('3/23', '상장', '오채영', '한독크린텍'),
    ('3/30', '공공', '구현경', '대전녹색환경지원센터'),
    ('3/30', '공공', '원창호', '한국항공우주연구원'),
    ('3/30', '공공', '백재현', '대전지방기상청'),
    ('3/30', '상장', '허준혁', '비플라이소프트'),
    ('3/30', '공공', '이승주', '한국수자원공사'),  # Keep this one
    ('3/30', '상장', '강체첵', 'kt&g'),
    ('3/30', '상장', '진승훈', '큐로셀'),
    ('3/23', '상장', '전은재', '쎄트렉아이'),
    ('4/6', '상장', '구현경', 'KT&G'),
    ('4/6', '공공', '오다인', '대전직업능력개발원'),
    ('4/6', '공공', '박해원', '정보통신기획평가원'),
    ('4/6', '상장', '이승주', '계룡건설산업'),
    ('4/6', '상장', '백재현', '에르코스'),
    ('4/6', '공공', '이현기', '한국철도시설공단'),
    ('4/13', '공공', '양하영', '국방과학연구소'),
    ('4/13', '공공', '이수호', '한국전자통신연구원'),
    ('4/13', '상장', '정현준', '동양에스텍'),
    ('4/13', '상장', '원창호', '펩트론'),
    ('4/13', '상장', '안은률', '대산f&b'),
    ('4/13', '공공', '이준성', '중소기업기술정보진흥원'),
    ('4/6', '공공', '유태수', '코레일테크(주)'), # Keep here, but will style it
    ('4/13', '상장', '곽현지', '골프존'),
    ('4/20', '공공', '구민경', '대전시민천문대'),
    ('4/20', '공공', '신재민', '한국가스기술공사'),
    ('4/20', '공공', '진승훈', '신용보증재단중앙회'),
    ('5/4', '공공', '김채은', '대전세관'),
    ('5/4', '상장', '오다인', '우성'),
    ('5/4', '상장', '이수호', '인텍플러스'),
    ('5/4', '상장', '박해원', '아이디스'),
    ('5/4', '상장', '이준희', 'KTcs'),
    ('5/4', '상장', '정구영', '코셈'),
    ('5/11', '공공', '안은률', '대전문화산업진흥협회'),
    ('5/11', '상장', '김주헌', '아이디스홀딩스'),
    ('5/11', '공공', '김주헌', '안전보건공단대전지역본부'),
    ('5/11', '상장', '이준성', '위드텍'),
    ('5/11', '상장', '유태수', '지노믹트리'),
    ('5/11', '공공', '한성경', '한국산림복지진흥원'),
    ('5/11', '공공', '이채원', '대전창조경제혁신센터'),
    ('5/18', '상장', '양하영', '알루코'),
    ('5/18', '공공', '정현준', '한국화학연구원'),
    ('5/18', '상장', '정하진', '원텍'),
    ('5/18', '공공', '정구영', 'k-water'),
    ('5/18', '상장', '이채원', '네오팜'),
    ('5/18', '공공', '전은재', '한국기계연구원'),
    ('5/18', '공공', '곽현지', '대전경제통상진흥원'),
    ('6/1', '상장', '구민경', '컨텍'),
    ('6/1', '공공', '양예지', '대전도시공사'),
    ('6/1', '상장', '신재민', '파이버프로'),
    ('6/1', '상장', '양예지', '레인보우로보틱스')
]

dates = ['3/23', '3/30', '4/6', '4/13', '4/20', '5/4', '5/11', '5/18', '6/1']

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "발표일정"

# Write headers
ws.append(["날짜", "공공기관", "상장기업"])
red_font = Font(color="FF0000")

for d in dates:
    pub_list = []
    pri_list = []
    
    for pd, pg, pn, pc in parsed_data:
        if pd == d:
            entry = f"{pn} ({pc})"
            if pn == '유태수' and pd == '4/6':
                entry += " (6주차(4/13) 게시판에 등록했으나 4/6으로 기재함)"
                
            if pg == '공공':
                pub_list.append(entry)
            else:
                pri_list.append(entry)
                
    pub_str = "\n".join(pub_list)
    pri_str = "\n".join(pri_list)
    
    ws.append([d, pub_str, pri_str])

# Adjust formatting
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=3):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='left' if cell.col_idx > 1 else 'center')

# Apply red font to Yoo Tae-soo
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=2):
    for cell in row:
        if cell.value and "유태수" in cell.value:
            # openpyxl doesn't support rich text (partial coloring) easily in this version.
            # We'll just color the entire cell red if it contains his name.
            cell.font = red_font
            
# Adjust column widths
ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 80
ws.column_dimensions['C'].width = 60

wb.save('c:/Users/Sumin/00_Projects/지역사회의이해발표시트/발표일정_세로형.xlsx')
